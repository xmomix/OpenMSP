# pyright: reportAttributeAccessIssue=false, reportIndexIssue=false, reportCallIssue=false, reportPossiblyUnboundVariable=false
# (Django 6 non pubblica py.typed: per pyright .objects e ._meta non esistono; gli stub di
#  openpyxl tipizzano Workbook.worksheets[0] come WriteOnlyWorksheet e requests.Response e'
#  assignata dentro un try. Rumore di tipi su codice gia' a terra, nessun comportamento coinvolto)
from django.shortcuts import render, redirect

from impostazioni.models import UtentiParametri
from impostazioni.models import ServiziParametri
from impostazioni.models import AnisServizi
from impostazioni.models import AnisParametri

from .utils import salva_log
from .utils import svuota_none
from .utils import converti_data
from .utils import cf_da_file
from .verifica_cf import verifica_cf

##from datetime import datetime, date
import datetime
from jose.constants import Algorithms
import http.client, urllib.parse
import hashlib
import random
import base64
import datetime
import uuid
import jwt
###import subprocess
import requests
import json
import io
import csv
import re
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from django.http import HttpResponse

# I payload di ANIS e ANIST non hanno un campo d'esito comune (IFS02/IFS03 portano solo
# l'elenco, ANIST porta frequentante/presenzaTitoli), quindi il giudizio va derivato dalla
# forma della risposta. Unica fonte di verita': la usano sia gli export sia la pagina unificata.
def _esito_verifica(dati):
    """None se la verifica ha prodotto dati, altrimenti il messaggio da mostrare."""
    if not isinstance(dati, dict):
        return "Verifica non eseguibile: risposta priva di esito"
    esito = dati.get("esito")
    if isinstance(esito, dict):
        return f"Errore tecnico ({esito.get('codice')}): {str(esito.get('descrizione'))[:180]}"
    for elenco in ("enrollments", "qualifications", "listaTitoli"):
        if elenco in dati:
            return None if dati.get(elenco) else "La richiesta effettuata non produce alcun risultato"
    if "frequentante" in dati:
        return None if dati.get("frequentante") else "Esito negativo della verifica"
    if "presenzaTitoli" in dati:
        return None if dati.get("presenzaTitoli") else "La richiesta effettuata non produce alcun risultato"
    return "Verifica non eseguibile: risposta priva di esito"


def _riga_errore(row, colonne):
    """Riga di esportazione per una verifica senza dati: [cf, messaggio, 'N/A'...], o [] se e' ok.
    Le righe in stringa ("CF Codice fiscale non corretto") escono []: le interpretano i rami
    else delle singole view di export."""
    dati = row[1] if isinstance(row, (list, tuple)) else row
    if not isinstance(dati, dict):
        return []
    messaggio = _esito_verifica(dati)
    if messaggio is None:
        return []
    if isinstance(row, (list, tuple)):
        cf = row[0]
    else:
        # ponytail: IFS02/IFS03 massivo appenda il solo payload, quindi qui il CF non c'e'.
        # Serve far appendere (cf, res) anche ai due view massivi IFS, poi cf = row[0] per tutti.
        cf = (dati.get("personal_data") or {}).get("tax_code", "N/A")
    return [cf, messaggio] + ["N/A"] * (colonne - 2)


ANIS_ESITO_FREQUENZA = {
    1: "Frequentante", 2: "Non Frequentante",
    3: "Frequentante su altro anno corso", 4: "Non piu Frequentante",
}


def _con_codice(nome, codice):
    return f"{nome} (cod. {codice})" if codice else (nome or "")


def _righe_servizio(id_caso, dati):
    """Righe della card: una per iscrizione / titolo / frequenza, come liste di coppie
    (etichetta, valore). Le coppie vuole cadono, il template non deve controllare nulla."""
    def ripulita(coppie):
        return [(etichetta, valore) for etichetta, valore in coppie if valore]
    if id_caso == 1:
        return [ripulita([
            ("Istituto", _con_codice(e.get("institute_name"), e.get("institute_code"))),
            ("Tipologia corso", _con_codice(e.get("programme_type_name"), e.get("programme_type_code"))),
            ("Nome del corso", _con_codice(e.get("degree_course_name"), e.get("degree_course_code"))),
            ("Classe", _con_codice(e.get("degree_class_name"), e.get("degree_class_code"))),
            ("Anno accademico", e.get("academic_year")),
            ("Anni durata corso", e.get("degree_course_year")),
        ]) for e in dati.get("enrollments") or []]
    if id_caso == 2:
        righe = []
        for qual in dati.get("qualifications") or []:
            voto = qual.get("qualification_grade_value")
            massimale = qual.get("qualification_grading_scale_maximum_grade")
            if voto == "QUALIFIED":
                voto = "Abilitato"
            elif voto == "110L":
                voto = f"110 cum laude su {massimale}"
            elif voto:
                voto = f"{voto} su {massimale}"
            righe.append(ripulita([
                ("Istituto", _con_codice(qual.get("institute_name"), qual.get("institute_code"))),
                ("Qualifica", qual.get("qualification_name")),
                ("Tipologia corso", _con_codice(qual.get("programme_type_name"), qual.get("programme_type_code"))),
                ("Nome del corso", _con_codice(qual.get("degree_course_name"), qual.get("degree_course_code"))),
                ("Classe", _con_codice(qual.get("degree_class_name"), qual.get("degree_class_code"))),
                ("Data conseguimento", qual.get("academic_qualification_date")),
                ("Valutazione", voto),
            ]))
        return righe
    if id_caso == 3:
        return [ripulita([
            ("Istituto principale", _con_codice(dati.get("denoIstitutoPrincipale"), dati.get("codiceIstitutoPrincipale"))),
            ("Plesso", _con_codice(dati.get("denominazionePlesso"), dati.get("codiceMeccanografico"))),
            ("Tipologia corso", dati.get("percorsoStudi")),
            ("Anno corso", dati.get("annoCorso")),
            ("Esito frequenza", ANIS_ESITO_FREQUENZA.get(dati.get("esitoFrequenza"), dati.get("esitoFrequenza"))),
        ])]
    righe = []
    for titolo in dati.get("listaTitoli") or []:
        votazione = titolo.get("votoFinale")
        if votazione and titolo.get("flagLode") == "S":
            votazione = f"{votazione} con lode"
        righe.append(ripulita([
            ("Titolo", _con_codice(titolo.get("denominazioneTitolo"), titolo.get("codiceTitolo"))),
            ("Istituto principale", _con_codice(titolo.get("denoIstitutoPrincipale"), titolo.get("codiceIstitutoPrincipale"))),
            ("Plesso", _con_codice(titolo.get("denominazionePlesso"), titolo.get("codiceMeccanografico"))),
            ("Votazione", votazione),
        ]))
    return righe


def anis_iscrizioni_export_excel(request):
    data = request.session.get("multi_data", [])  # Oppure recuperalo come preferisci
    wb = Workbook()
    ws = wb.worksheets[0]   # openpyxl: stessa cosa di .active su un workbook nuovo, ma senza Optional
    ws.append(["Codice Fiscale", "Istituto", "Tipologia corso", "Nome corso", "Classe", "Anno accademico", "Durata corso"])

    # Dizionario per la larghezza massima di ogni colonna
    max_lengths = [len(str(cell.value or "")) for cell in ws[1]]  # Larghezze iniziali dall'intestazione

    if not data:
        return HttpResponse("Nessun dato disponibile per l'esportazione", status=400)

    for i, row in enumerate(data, 1):
        color = "FFFFFF"  # default bianco

        err = _riga_errore(row, 7)
        if err:
            color = "FFC7CE"
            ws.append(err)
        elif isinstance(row, dict):
            cf = row['personal_data']['tax_code']
            if len(row.get("enrollments")) == 0 :
                color = "FFC7CE"
                ws.append([cf, "La richiesta effettuata non produce alcun risultato", "N/A", "N/A", "N/A", "N/A", "N/A"])
            else:
                for enrol in row['enrollments']:
                    color = "C6EFCE"  # verde chiaro
                    istituto = enrol['institute_name'] + " (" + enrol['institute_code'] + ")"
                    tipo_corso = enrol['programme_type_code']
                    nome_corso = enrol['degree_course_code']
                    classe = enrol['degree_class_code']
                    anno_accademico = enrol['academic_year']
                    durata_corso = enrol['degree_course_year']
                    ws.append([cf, istituto, tipo_corso, nome_corso, classe, anno_accademico, durata_corso])
        else:
            split_row = row.split()
            cf = split_row[0]
            if "Codice" in split_row[1]:
                color = "FFC7CE"  # rosso
                ws.append([cf, "Codice fiscale non corretto", "N/A", "N/A", "N/A", "N/A", "N/A"])

        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        for col in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row):
            for i, cell in enumerate(col):
                cell.fill = fill
                value_length = len(str(cell.value)) if cell.value else 0
                if len(max_lengths) <= i:
                    max_lengths.append(value_length)
                else:
                    max_lengths[i] = max(max_lengths[i], value_length)

    # Imposta larghezza colonne
    for i, width in enumerate(max_lengths, 1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    # Scrive l'Excel su un buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=EsitoAnisIscrizioni.xlsx'
    return response


def anis_iscrizioni_export_csv(request):
    data = request.session.get("multi_data", [])
    
    response = HttpResponse(
        content_type='text/csv',
        headers={'Content-Disposition': 'attachment; filename="EsitoAnisIscrizioni.csv"'},
    )

    writer = csv.writer(response, delimiter=';')
    writer.writerow(["Codice Fiscale", "Istituto", "Tipologia corso", "Nome corso", "Classe", "Anno accademico", "Durata corso"])

    if not data:
        return HttpResponse("Nessun dato disponibile per l'esportazione", status=400)

    for row in data:
        err = _riga_errore(row, 7)
        if err:
            writer.writerow(err)
        elif isinstance(row, dict):
            cf = row['personal_data']['tax_code']
            if len(row.get("enrollments")) == 0:
                writer.writerow([cf, "La richiesta effettuata non produce alcun risultato", "N/A", "N/A", "N/A", "N/A", "N/A"])
            else:
                for enrol in row['enrollments']:
                    istituto = enrol['institute_name'] + " (" + enrol['institute_code'] + ")"
                    tipo_corso = enrol['programme_type_code']
                    nome_corso = enrol['degree_course_code']
                    classe = enrol['degree_class_code']
                    anno_accademico = enrol['academic_year']
                    durata_corso = enrol['degree_course_year']
                    writer.writerow([cf, istituto, tipo_corso, nome_corso, classe, anno_accademico, durata_corso])
        else:
            split_row = row.split()
            cf = split_row[0]
            if "Codice" in split_row[1]:
                writer.writerow([cf, "Codice fiscale non corretto", "N/A", "N/A", "N/A", "N/A", "N/A"])

    return response


def anis_titoli_export_excel(request):
    data = request.session.get("multi_data", [])  # Oppure recuperalo come preferisci
    wb = Workbook()
    ws = wb.worksheets[0]   # openpyxl: come .active su un workbook nuovo, ma senza Optional
    ws.append(["Codice Fiscale", "Istituto", "Qualifica", "Tipologia corso", "Nome corso", "Classe", "Data conseguimento", "Valutazione"])

    # Dizionario per la larghezza massima di ogni colonna
    max_lengths = [len(str(cell.value or "")) for cell in ws[1]]  # Larghezze iniziali dall'intestazione

    if not data:
        return HttpResponse("Nessun dato disponibile per l'esportazione", status=400)

    for i, row in enumerate(data, 1):
        color = "FFFFFF"  # default bianco

        err = _riga_errore(row, 8)
        if err:
            color = "FFC7CE"
            ws.append(err)
        elif isinstance(row, dict):
            cf = row['personal_data']['tax_code']
            if len(row.get("qualifications")) == 0 :
                color = "FFC7CE"
                ws.append([cf, "La richiesta effettuata non produce alcun risultato", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A"])
                fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                for col in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row):
                    for i, cell in enumerate(col):
                        cell.fill = fill
                        value_length = len(str(cell.value)) if cell.value else 0
                        if len(max_lengths) <= i:
                            max_lengths.append(value_length)
                        else:
                            max_lengths[i] = max(max_lengths[i], value_length)
            else:
                for qual in row['qualifications']:
                    color = "C6EFCE"  # verde chiaro
                    istituto = qual['institute_name'] + " (" + qual['institute_code'] + ")"
                    qualifica = qual['qualification_name']
                    tipo_corso = qual['programme_type_code']
                    nome_corso = qual['degree_course_code']
                    classe = qual['degree_class_code']
                    data_conseg = qual['academic_qualification_date']
                    valutazione = qual['qualification_grade_value']
                    if valutazione == "QUALIFIED" or valutazione == "Abilitato":
                        valutazione = "Abilitato"
                    else:
                        if valutazione == "110L":
                            valutazione = "110 cum laude"
                        valutazione = valutazione + " su " + qual['qualification_grading_scale_maximum_grade']
                    ws.append([cf, istituto, qualifica, tipo_corso, nome_corso, classe, data_conseg, valutazione])
                    
                    fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                    for col in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row):
                        for i, cell in enumerate(col):
                            cell.fill = fill
                            value_length = len(str(cell.value)) if cell.value else 0
                            if len(max_lengths) <= i:
                                max_lengths.append(value_length)
                            else:
                                max_lengths[i] = max(max_lengths[i], value_length)
        else:
            split_row = row.split()
            cf = split_row[0]
            if "Codice" in split_row[1]:
                color = "FFC7CE"  # rosso
                ws.append([cf, "Codice fiscale non corretto", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A"])
                
                fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                for col in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row):
                    for i, cell in enumerate(col):
                        cell.fill = fill
                        value_length = len(str(cell.value)) if cell.value else 0
                        if len(max_lengths) <= i:
                            max_lengths.append(value_length)
                        else:
                            max_lengths[i] = max(max_lengths[i], value_length)

    # Imposta larghezza colonne
    for i, width in enumerate(max_lengths, 1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    # Scrive l'Excel su un buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=EsitoAnisTitoli.xlsx'
    return response


def anist_frequenze_export_excel(request):
    data = request.session.get("multi_data", [])  # Oppure recuperalo come preferisci
    wb = Workbook()
    ws = wb.worksheets[0]   # openpyxl: come .active su un workbook nuovo, ma senza Optional
    ws.append(["Codice Fiscale", "Istituto principale", "Plesso", "Tipologia corso", "Anno corso", "Esito frequenza"])

    # Dizionario per la larghezza massima di ogni colonna
    max_lengths = [len(str(cell.value or "")) for cell in ws[1]]  # Larghezze iniziali dall'intestazione

    if not data:
        return HttpResponse("Nessun dato disponibile per l'esportazione", status=400)

    for i, row in enumerate(data, 1):
        color = "FFFFFF"  # default bianco

        err = _riga_errore(row, 6)
        if err:
            color = "FFC7CE"
            ws.append(err)
        elif isinstance(row, list):
            cf = row[0]
            if isinstance(row[1], dict):
                if row[1]['frequentante'] == False :
                    color = "FFC7CE"
                    ws.append([cf, "La richiesta effettuata non produce alcun risultato", "N/A", "N/A", "N/A", "N/A"])
                else:
                    color = "C6EFCE"  # verde chiaro
                    istituto = row[1]['denoIstitutoPrincipale'] + " (" + row[1]['codiceIstitutoPrincipale'] + ")"
                    plesso = row[1]['denominazionePlesso'] + " (" + row[1]['codiceMeccanografico'] + ")"
                    tipo_corso = row[1]['percorsoStudi']
                    anno_accademico = row[1]['annoCorso']
                    esito_map = {1: "Frequentante", 2: "Non Frequentante", 3: "Frequentante su altro anno corso", 4: "Non più Frequentante"}
                    esito = esito_map.get(row[1].get('esitoFrequenza'), row[1].get('esitoFrequenza'))
                    ws.append([cf, istituto, plesso, tipo_corso, anno_accademico, esito])
            else:
                color = "FFC7CE"
                ws.append([cf, row[1], "N/A", "N/A", "N/A", "N/A"])
        else:
            split_row = row.split()
            cf = split_row[0]
            if "Codice" in split_row[1]:
                color = "FFC7CE"  # rosso
                ws.append([cf, "Codice fiscale non corretto", "N/A", "N/A", "N/A", "N/A"])

        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        for col in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row):
            for i, cell in enumerate(col):
                cell.fill = fill
                value_length = len(str(cell.value)) if cell.value else 0
                if len(max_lengths) <= i:
                    max_lengths.append(value_length)
                else:
                    max_lengths[i] = max(max_lengths[i], value_length)

    # Imposta larghezza colonne
    for i, width in enumerate(max_lengths, 1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    # Scrive l'Excel su un buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=EsitoAnistFrequenze.xlsx'
    return response


def anist_frequenze_export_csv(request):
    data = request.session.get("multi_data", [])
    
    response = HttpResponse(
        content_type='text/csv',
        headers={'Content-Disposition': 'attachment; filename="EsitoAnistFrequenze.csv"'},
    )

    writer = csv.writer(response, delimiter=';')
    writer.writerow(["Codice Fiscale", "Istituto principale", "Plesso", "Tipologia corso", "Anno corso", "Esito frequenza"])

    if not data:
        return HttpResponse("Nessun dato disponibile per l'esportazione", status=400)

    for row in data:
        err = _riga_errore(row, 6)
        if err:
            writer.writerow(err)
        elif isinstance(row, list):
            cf = row[0]
            if isinstance(row[1], dict):
                if row[1]['frequentante'] == False:
                    writer.writerow([cf, "La richiesta effettuata non produce alcun risultato", "N/A", "N/A", "N/A", "N/A"])
                else:
                    istituto = row[1]['denoIstitutoPrincipale'] + " (" + row[1]['codiceIstitutoPrincipale'] + ")"
                    plesso = row[1]['denominazionePlesso'] + " (" + row[1]['codiceMeccanografico'] + ")"
                    tipo_corso = row[1]['percorsoStudi']
                    anno_accademico = row[1]['annoCorso']
                    esito_map = {1: "Frequentante", 2: "Non Frequentante", 3: "Frequentante su altro anno corso", 4: "Non più Frequentante"}
                    esito = esito_map.get(row[1].get('esitoFrequenza'), row[1].get('esitoFrequenza'))
                    writer.writerow([cf, istituto, plesso, tipo_corso, anno_accademico, esito])
            else:
                writer.writerow([cf, row[1], "N/A", "N/A", "N/A", "N/A"])
        else:
            split_row = row.split()
            cf = split_row[0]
            if "Codice" in split_row[1]:
                writer.writerow([cf, "Codice fiscale non corretto", "N/A", "N/A", "N/A", "N/A"])

    return response


def anist_titoli_export_excel(request):
    data = request.session.get("multi_data", [])  # Oppure recuperalo come preferisci
    wb = Workbook()
    ws = wb.worksheets[0]   # openpyxl: come .active su un workbook nuovo, ma senza Optional
    ws.append(["Codice Fiscale", "Titolo", "Istituto principale", "Plesso", "Votazione"])

    # Dizionario per la larghezza massima di ogni colonna
    max_lengths = [len(str(cell.value or "")) for cell in ws[1]]  # Larghezze iniziali dall'intestazione

    if not data:
        return HttpResponse("Nessun dato disponibile per l'esportazione", status=400)

    for i, row in enumerate(data, 1):
        color = "FFFFFF"  # default bianco

        err = _riga_errore(row, 5)
        if err:
            color = "FFC7CE"
            ws.append(err)
        elif isinstance(row, list):
            cf = row[0]
            if row[1]['presenzaTitoli'] == False:
                color = "FFC7CE"
                valori = [cf, "La richiesta effettuata non produce alcun risultato", "N/A", "N/A", "N/A"]
                ws.append(valori)
            else:
                for titol in row[1]['listaTitoli']:
                    color = "C6EFCE"
                    titolo = titol['denominazioneTitolo']
                    istituto = f"{titol['denoIstitutoPrincipale']} ({titol['codiceIstitutoPrincipale']})"
                    plesso = f"{titol['denominazionePlesso']} ({titol['codiceMeccanografico']})"
                    valutazione = titol['votoFinale']
                    if titol['flagLode'] == "S":
                        valutazione += " con lode"
                    valori = [cf, titolo, istituto, plesso, valutazione]
                    ws.append(valori)
        else:
            split_row = row.split()
            cf = split_row[0]
            if "Codice" in split_row[1]:
                color = "FFC7CE"
                valori = [cf, "Codice fiscale non corretto", "N/A", "N/A", "N/A"]
                ws.append(valori)

        # Applica colore all'ultima riga scritta
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        for col in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row):
            for i, cell in enumerate(col):
                cell.fill = fill
                value_length = len(str(cell.value)) if cell.value else 0
                if len(max_lengths) <= i:
                    max_lengths.append(value_length)
                else:
                    max_lengths[i] = max(max_lengths[i], value_length)

    # Imposta larghezza colonne
    for i, width in enumerate(max_lengths, 1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    # Scrive l'Excel su un buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=EsitoAnistTitoli.xlsx'
    return response


def anist_titoli_export_csv(request):
    data = request.session.get("multi_data", [])
    
    response = HttpResponse(
        content_type='text/csv',
        headers={'Content-Disposition': 'attachment; filename="EsitoAnistTitoli.csv"'},
    )

    writer = csv.writer(response, delimiter=';')
    writer.writerow(["Codice Fiscale", "Titolo", "Istituto principale", "Plesso", "Votazione"])

    if not data:
        return HttpResponse("Nessun dato disponibile per l'esportazione", status=400)

    for row in data:
        err = _riga_errore(row, 5)
        if err:
            writer.writerow(err)
        elif isinstance(row, list):
            cf = row[0]
            if row[1]['presenzaTitoli'] == False:
                writer.writerow([cf, "La richiesta effettuata non produce alcun risultato", "N/A", "N/A", "N/A"])
            else:
                for titol in row[1]['listaTitoli']:
                    titolo = titol['denominazioneTitolo']
                    istituto = f"{titol['denoIstitutoPrincipale']} ({titol['codiceIstitutoPrincipale']})"
                    plesso = f"{titol['denominazionePlesso']} ({titol['codiceMeccanografico']})"
                    valutazione = titol['votoFinale']
                    if titol['flagLode'] == "S":
                        valutazione += " con lode"
                    writer.writerow([cf, titolo, istituto, plesso, valutazione])
        else:
            split_row = row.split()
            cf = split_row[0]
            if "Codice" in split_row[1]:
                writer.writerow([cf, "Codice fiscale non corretto", "N/A", "N/A", "N/A"])

    return response


def anis_iscrizioni_singola(request):
    if request.user.id:
        utente_sessione = UtentiParametri.objects.get(id=request.user.id)
        utente_abilitato = utente_sessione.anis_IFS02_singolo
        if request.method == 'POST':
            data = []
            cf = request.POST.get('input_CF')
            data.append(cf)
            correttezza_cf = verifica_cf(cf)
            if correttezza_cf == 1 or correttezza_cf == 2:
                res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, cf, 1)
                data.append(res)
                data = converti_data(data)
                salva_log(request.user, "Verifica ANIS - IFS02 - Iscrizioni Singolo", "Verificato utente " + cf, purposeid=purp_id, resp_status=status, token_id=tok_id)
            else:
                data.append(str(cf) + " Codice fiscale non corretto")
                salva_log(request.user, "Verifica ANIS - IFS02 - Iscrizioni Singolo", "Verificato utente " + cf)
            return render(request, 'anis_iscrizioni_singola.html', {'data': data, 'utente_abilitato': utente_abilitato })
    else:
        utente_abilitato = False

    return render(request, 'anis_iscrizioni_singola.html', { 'utente_abilitato': utente_abilitato })


def anis_iscrizioni_massiva(request):
    if request.user.id:
        utente_sessione = UtentiParametri.objects.get(id=request.user.id)
        utente_abilitato = utente_sessione.anis_IFS02_massivo
        if request.method == 'POST':
            csv_file = request.FILES['cf_csv']
            data = []
            contatore = 0
            if csv_file.name.endswith('.csv'):
                csv_file_text = io.TextIOWrapper(csv_file.file, encoding='utf-8')
                csv_reader = csv.reader(csv_file_text)
                for row in csv_reader:
                    if row[0]:  # Se row[0] non è vuoto o None
                        correttezza_cf = verifica_cf(row[0].strip().upper())
                        if correttezza_cf == 1 or correttezza_cf == 2:
                            res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, row[0].strip().upper(), 1)
                            data.append(res)
                            salva_log(request.user,"Verifica ANIS - IFS02 - Iscrizioni Massivo", "Verificato utente " + row[0].strip().upper(), purposeid=purp_id, resp_status=status, token_id=tok_id)
                        else:
                            data.append( str(row[0].strip().upper()) + " Codice fiscale non corretto")
                            salva_log(request.user,"Verifica ANIS - IFS02 - Iscrizioni Massivo", "Verificato utente " + row[0].strip().upper())
                        contatore += 1
                data = converti_data(data)
                request.session["multi_data"] = data  # <--- Salva i dati nella sessione
                return render(request, 'anis_iscrizioni_massiva.html', {'data': data, 'utente_abilitato': utente_abilitato })
            elif csv_file.name.endswith('.xlsx') or csv_file.name.endswith('.XLSX') or csv_file.name.endswith('.Xlsx'):
                wb = openpyxl.load_workbook(csv_file)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=1, values_only=True):
                    if row[0]:
                        correttezza_cf = verifica_cf(row[0].strip().upper())
                        if correttezza_cf == 1 or correttezza_cf == 2:
                            res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, row[0].strip().upper(), 1)
                            data.append(res)
                            salva_log(request.user,"Verifica ANIS - IFS02 - Iscrizioni Massivo", "Verificato utente " + row[0].strip().upper(), purposeid=purp_id, resp_status=status, token_id=tok_id)
                        else:
                            data.append( str(row[0].strip().upper()) + " Codice fiscale non corretto")
                            salva_log(request.user,"Verifica ANIS - IFS02 - Iscrizioni Massivo", "Verificato utente " + row[0].strip().upper())
                        contatore += 1
                data = converti_data(data)
                request.session["multi_data"] = data  # <--- Salva i dati nella sessione
                return render(request, 'anis_iscrizioni_massiva.html', {'data': data, 'utente_abilitato': utente_abilitato })

            else:
                salva_log(request.user,"Verifica ANIS - IFS02 - Iscrizioni Massivo", "Errore caricamento file CSV")
                return render(request, 'anis_iscrizioni_massiva.html', {'error': 'Il file non è un CSV', 'utente_abilitato': utente_abilitato })
    else:
        utente_abilitato = False
    return render(request, 'anis_iscrizioni_massiva.html' , { 'utente_abilitato': utente_abilitato })


def anis_titoli_singola(request):
    if request.user.id:
        utente_sessione = UtentiParametri.objects.get(id=request.user.id)
        utente_abilitato = utente_sessione.anis_IFS03_singolo
        if request.method == 'POST':
            data = []
            cf = request.POST.get('input_CF')
            data.append(cf)
            correttezza_cf = verifica_cf(cf)
            if correttezza_cf == 1 or correttezza_cf == 2:
                res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, cf, 2)
                data.append(res)
                data = converti_data(data)
                salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Singolo", "Verificato utente " + cf, purposeid=purp_id, resp_status=status, token_id=tok_id)
            else:
                data.append(str(cf) + " Codice fiscale non corretto")
                salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Singolo", "Verificato utente " + cf)

            return render(request, 'anis_titoli_singola.html', {'data': data, 'utente_abilitato': utente_abilitato })
    else:
        utente_abilitato = False
    return render(request, 'anis_titoli_singola.html' , { 'utente_abilitato': utente_abilitato })


def anis_titoli_massiva(request):
    if request.user.id:
        utente_sessione = UtentiParametri.objects.get(id=request.user.id)
        utente_abilitato = utente_sessione.anis_IFS03_massivo
        if request.method == 'POST':
            csv_file = request.FILES['cf_csv']
            data = []
            contatore = 0
            if csv_file.name.endswith('.csv'):
                csv_file_text = io.TextIOWrapper(csv_file.file, encoding='utf-8')
                csv_reader = csv.reader(csv_file_text)
                for row in csv_reader:
                    if row[0]:  # Se row[0] non è vuoto o None
                        correttezza_cf = verifica_cf(row[0].strip().upper())
                        if correttezza_cf == 1 or correttezza_cf == 2:
                            res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, row[0].strip().upper(), 2)
                            data.append(res)
                            salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Massivo", "Verificato utente " + row[0].strip().upper(), purposeid=purp_id, resp_status=status, token_id=tok_id)
                        else:
                            data.append( str(row[0].strip().upper()) + " Codice fiscale non corretto")
                            salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Massivo", "Verificato utente " + row[0].strip().upper())
                        contatore += 1
                salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Massivo", "Verificati n. " + str(contatore) + " CF")
                # Modifica QUALIFIED in Abilitato
                for item in data:
                    if isinstance(item, dict) and 'qualifications' in item:
                        for qual in item['qualifications']:
                            if qual.get('qualification_grade_value') == 'QUALIFIED':
                                qual['qualification_grade_value'] = 'Abilitato'
                data = converti_data(data)
                request.session["multi_data"] = data  # <--- Salva i dati nella sessione
                return render(request, 'anis_titoli_massiva.html', {'data': data, 'utente_abilitato': utente_abilitato })

            elif csv_file.name.endswith('.xlsx') or csv_file.name.endswith('.XLSX') or csv_file.name.endswith('.Xlsx'):
                wb = openpyxl.load_workbook(csv_file)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=1, values_only=True):
                    if row[0]:
                        correttezza_cf = verifica_cf(row[0].strip().upper())
                        if correttezza_cf == 1 or correttezza_cf == 2:
                            res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, row[0].strip().upper(), 2)
                            data.append(res)
                            salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Massivo", "Verificato utente " + row[0].strip().upper(), purposeid=purp_id, resp_status=status, token_id=tok_id)
                        else:
                            data.append( str(row[0].strip().upper()) + " Codice fiscale non corretto")
                            salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Massivo", "Verificato utente " + row[0].strip().upper())
                        contatore += 1
                salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Massivo", "Verificati n. " + str(contatore) + " CF")
                # Modifica QUALIFIED in Abilitato
                for item in data:
                    if isinstance(item, dict) and 'qualifications' in item:
                        for qual in item['qualifications']:
                            if qual.get('qualification_grade_value') == 'QUALIFIED':
                                qual['qualification_grade_value'] = 'Abilitato'
                data = converti_data(data)
                request.session["multi_data"] = data  # <--- Salva i dati nella sessione
                return render(request, 'anis_titoli_massiva.html', {'data': data, 'utente_abilitato': utente_abilitato })

            else:
                salva_log(request.user,"Verifica ANIS - IFS03 - Titoli Massivo", "Errore caricamento file CSV")
                return render(request, 'anis_titoli_massiva.html', {'error': 'Il file non è un CSV', 'utente_abilitato': utente_abilitato })
    else:
        utente_abilitato = False
    return render(request, 'anis_titoli_massiva.html' , { 'utente_abilitato': utente_abilitato })



def anist_frequenze_singola(request):
    if request.user.id:
        utente_sessione = UtentiParametri.objects.get(id=request.user.id)
        utente_abilitato = utente_sessione.anist_frequenze_singolo
        if request.method == 'POST':
            data = []
            cf = request.POST.get('input_CF')
            data.append(cf)
            correttezza_cf = verifica_cf(cf)
            if correttezza_cf == 1 or correttezza_cf == 2:
                res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, cf, 3)
                data.append(res)
                data = converti_data(data)
                salva_log(request.user,"Verifica ANIST - Frequenze Singolo", "Verificato utente " + cf, purposeid=purp_id, resp_status=status, token_id=tok_id)
            else:
                data.append(str(cf) + " Codice fiscale non corretto")
                salva_log(request.user,"Verifica ANIST - Frequenze Singolo", "Verificato utente " + cf)

            return render(request, 'anist_frequenze_singola.html', {'data': data, 'utente_abilitato': utente_abilitato })
    else:
        utente_abilitato = False

    return render(request, 'anist_frequenze_singola.html', { 'utente_abilitato': utente_abilitato })


def anist_frequenze_massiva(request):
    if request.user.id:
        utente_sessione = UtentiParametri.objects.get(id=request.user.id)
        utente_abilitato = utente_sessione.anist_frequenze_massivo
        if request.method == 'POST':
            csv_file = request.FILES['cf_csv']
            data = []
            contatore = 0
            if csv_file.name.endswith('.csv'):
                csv_file_text = io.TextIOWrapper(csv_file.file, encoding='utf-8')
                csv_reader = csv.reader(csv_file_text)
                for row in csv_reader:
                    if row[0]:  # Se row[0] non è vuoto o None
                        correttezza_cf = verifica_cf(row[0].strip().upper())
                        if correttezza_cf == 1 or correttezza_cf == 2:
                            res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, row[0].strip().upper(), 3)
                            data.append((row[0].strip().upper(), res))
                            salva_log(request.user,"Verifica ANIST - Frequenze Massivo", "Verificato utente " + row[0].strip().upper(), purposeid=purp_id, resp_status=status, token_id=tok_id)
                        else:
                            data.append( str(row[0].strip().upper()) + " Codice fiscale non corretto")
                            salva_log(request.user,"Verifica ANIST - Frequenze Massivo", "Verificato utente " + row[0].strip().upper())
                        contatore += 1
                data = converti_data(data)
                request.session["multi_data"] = data  # <--- Salva i dati nella sessione
                return render(request, 'anist_frequenze_massiva.html', {'data': data, 'utente_abilitato': utente_abilitato })
            elif csv_file.name.endswith('.xlsx') or csv_file.name.endswith('.XLSX') or csv_file.name.endswith('.Xlsx'):
                wb = openpyxl.load_workbook(csv_file)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=1, values_only=True):
                    if row[0]:
                        correttezza_cf = verifica_cf(row[0].strip().upper())
                        if correttezza_cf == 1 or correttezza_cf == 2:
                            res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, row[0].strip().upper(), 3)
                            data.append((row[0].strip().upper(), res))
                            salva_log(request.user,"Verifica ANIST - Frequenze Massivo", "Verificato utente " + row[0].strip().upper(), purposeid=purp_id, resp_status=status, token_id=tok_id)
                        else:
                            data.append(str(row[0].strip().upper()) + " Codice fiscale non corretto")
                            salva_log(request.user,"Verifica ANIST - Frequenze Massivo", "Verificato utente " + row[0].strip().upper())
                        contatore += 1
                data = converti_data(data)
                request.session["multi_data"] = data  # <--- Salva i dati nella sessione
                return render(request, 'anist_frequenze_massiva.html', {'data': data, 'utente_abilitato': utente_abilitato })
            else:
                salva_log(request.user,"Verifica ANIST - Frequenze Massivo", "Errore caricamento file CSV")
                return render(request, 'anist_frequenze_massiva.html', {'error': 'Il file non è un CSV', 'utente_abilitato': utente_abilitato })
    else:
        utente_abilitato = False
    return render(request, 'anist_frequenze_massiva.html', { 'utente_abilitato': utente_abilitato })


def anist_titoli_singola(request):
    if request.user.id:
        utente_sessione = UtentiParametri.objects.get(id=request.user.id)
        utente_abilitato = utente_sessione.anist_titoli_singolo
        if request.method == 'POST':
            data = []
            cf = request.POST.get('input_CF')
            data.append(cf)
            correttezza_cf = verifica_cf(cf)
            if correttezza_cf == 1 or correttezza_cf == 2:
                res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, cf, 4)
                data.append(res)
                data = converti_data(data)
                salva_log(request.user,"Verifica ANIST - Titoli Singolo", "Verificato utente " + cf, purposeid=purp_id, resp_status=status, token_id=tok_id)
            else:
                data.append(str(cf) + " Codice fiscale non corretto")
                salva_log(request.user,"Verifica ANIST - Titoli Singolo", "Verificato utente " + cf)

            return render(request, 'anist_titoli_singola.html', {'data': data, 'utente_abilitato': utente_abilitato })
    else:
        utente_abilitato = False

    return render(request, 'anist_titoli_singola.html', { 'utente_abilitato': utente_abilitato })


def anist_titoli_massiva(request):
    if request.user.id:
        utente_sessione = UtentiParametri.objects.get(id=request.user.id)
        utente_abilitato = utente_sessione.anist_titoli_massivo
        if request.method == 'POST':
            csv_file = request.FILES['cf_csv']
            data = []
            contatore = 0
            if csv_file.name.endswith('.csv'):
                csv_file_text = io.TextIOWrapper(csv_file.file, encoding='utf-8')
                csv_reader = csv.reader(csv_file_text)
                for row in csv_reader:
                    if row[0]:  # Se row[0] non è vuoto o None
                        correttezza_cf = verifica_cf(row[0].strip().upper())
                        if correttezza_cf == 1 or correttezza_cf == 2:
                            res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, row[0].strip().upper(), 4)
                            data.append((row[0].strip().upper(), res))
                            salva_log(request.user,"Verifica ANIST - Titoli Massivo", "Verificato utente " + row[0].strip().upper(), purposeid=purp_id, resp_status=status, token_id=tok_id)
                        else:
                            data.append( str(row[0].strip().upper()) + " Codice fiscale non corretto")
                            salva_log(request.user,"Verifica ANIST - Titoli Massivo", "Verificato utente " + row[0].strip().upper())
                        contatore += 1
                data = converti_data(data)
                request.session["multi_data"] = data  # <--- Salva i dati nella sessione
                return render(request, 'anist_titoli_massiva.html', {'data': data, 'utente_abilitato': utente_abilitato })
            elif csv_file.name.endswith('.xlsx') or csv_file.name.endswith('.XLSX') or csv_file.name.endswith('.Xlsx'):
                wb = openpyxl.load_workbook(csv_file)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=1, values_only=True):
                    if row[0]:
                        correttezza_cf = verifica_cf(row[0].strip().upper())
                        if correttezza_cf == 1 or correttezza_cf == 2:
                            res, status, purp_id, tok_id = anis_verifica_utente(request.user.username, row[0].strip().upper(), 4)
                            data.append((row[0].strip().upper(), res))
                            salva_log(request.user,"Verifica ANIST - Titoli Massivo", "Verificato utente " + row[0].strip().upper(), purposeid=purp_id, resp_status=status, token_id=tok_id)
                        else:
                            data.append(str(row[0].strip().upper()) + " Codice fiscale non corretto")
                            salva_log(request.user,"Verifica ANIST - Titoli Massivo", "Verificato utente " + row[0].strip().upper())
                        contatore += 1
                data = converti_data(data)
                request.session["multi_data"] = data  # <--- Salva i dati nella sessione
                return render(request, 'anist_titoli_massiva.html', {'data': data, 'utente_abilitato': utente_abilitato })
            else:
                salva_log(request.user,"Verifica ANIST - Titoli Massivo", "Errore caricamento file CSV")
                return render(request, 'anist_titoli_massiva.html', {'error': 'Il file non è un CSV', 'utente_abilitato': utente_abilitato })
    else:
        utente_abilitato = False
    return render(request, 'anist_titoli_massiva.html', { 'utente_abilitato': utente_abilitato })


# Una riga per ogni interrogazione del menu "Istruzione": id_caso e forma della risposta
# stanno gia' in AnisParametri/anis_verifica_utente, qui serve l'abbinamento con il doppio
# gate (ServiziParametri.attivo AND il flag su utenti_parametri), con i nomi di audit delle
# viste singole (la pagina dei log non si biforca) e con gli export gia' esistenti.
ANIS_SERVIZI = {
    'IFS02': {
        'id_caso': 1, 'codice': 'anis_IFS02', 'label': 'ANIS - Iscrizioni universitarie',
        'titolo': 'Iscrizioni', 'con_cf': False,
        'singolo': 'anis_IFS02_singolo', 'massivo': 'anis_IFS02_massivo',
        'log': {'singola': 'Verifica ANIS - IFS02 - Iscrizioni Singolo',
                'massiva': 'Verifica ANIS - IFS02 - Iscrizioni Massivo'},
        'export': {'excel': 'anis_iscrizioni_export_excel', 'csv': 'anis_iscrizioni_export_csv'},
    },
    'IFS03': {
        'id_caso': 2, 'codice': 'anis_IFS03', 'label': 'ANIS - Titoli di studio universitari',
        'titolo': 'Titoli conseguiti', 'con_cf': False,
        'singolo': 'anis_IFS03_singolo', 'massivo': 'anis_IFS03_massivo',
        'log': {'singola': 'Verifica ANIS - IFS03 - Titoli Singolo',
                'massiva': 'Verifica ANIS - IFS03 - Titoli Massivo'},
        'export': {'excel': 'anis_titoli_export_excel'},
    },
    'ANIST_FREQ': {
        'id_caso': 3, 'codice': 'anist_frequenze', 'label': 'ANIST - Frequenza scolastica',
        'titolo': 'Frequenza di', 'con_cf': True,
        'singolo': 'anist_frequenze_singolo', 'massivo': 'anist_frequenze_massivo',
        'log': {'singola': 'Verifica ANIST - Frequenze Singolo',
                'massiva': 'Verifica ANIST - Frequenze Massivo'},
        'export': {'excel': 'anist_frequenze_export_excel', 'csv': 'anist_frequenze_export_csv'},
    },
    'ANIST_TITOLI': {
        'id_caso': 4, 'codice': 'anist_titoli', 'label': 'ANIST - Titoli di studio',
        'titolo': 'Titoli conseguiti', 'con_cf': True,
        'singolo': 'anist_titoli_singolo', 'massivo': 'anist_titoli_massivo',
        'log': {'singola': 'Verifica ANIST - Titoli Singolo',
                'massiva': 'Verifica ANIST - Titoli Massivo'},
        'export': {'excel': 'anist_titoli_export_excel', 'csv': 'anist_titoli_export_csv'},
    },
}


def istruzione(request):
    """Pagina unificata ANIS/ANIST: un form, quattro interrogazioni, singola o massiva.
    Il protocollo AgID resta in anis_verifica_utente, qui cambiano solo id_caso e i permessi."""
    permessi = {}
    servizi = []
    utente_abilitato = False
    if request.user.is_authenticated:
        utente = UtentiParametri.objects.filter(id=request.user.id).first()
        attivi = {servizio.codice_servizio: servizio.attivo for servizio in ServiziParametri.objects.all()}
        for chiave, conf in ANIS_SERVIZI.items():
            abilitato = attivi.get(conf['codice'], False) and bool(utente)
            permessi[chiave] = {
                'singola': abilitato and bool(getattr(utente, conf['singolo'], False)),
                'massiva': abilitato and bool(getattr(utente, conf['massivo'], False)),
            }
            servizi.append({'chiave': chiave, 'label': conf['label'],
                            'singola': permessi[chiave]['singola'],
                            'massiva': permessi[chiave]['massiva']})
            utente_abilitato = utente_abilitato or any(permessi[chiave].values())

    risultati = []
    interrogazione = {'titolo': '', 'modalita': '', 'export': {}}
    servizio_scelto = ''
    modalita_scelta = ''
    error = None

    if request.method == 'POST':
        servizio_scelto = request.POST.get('servizio_istruzione') or ''
        modalita_scelta = request.POST.get('modalita') or ''
        conf = ANIS_SERVIZI.get(servizio_scelto)
        if conf is None or modalita_scelta not in ('singola', 'massiva') \
                or not permessi.get(servizio_scelto, {}).get(modalita_scelta, False):
            error = 'Non sei abilitato per questa tipologia di ricerca.'
        else:
            if modalita_scelta == 'singola':
                cf_lista = [request.POST.get('input_CF', '').strip().upper()]
            else:
                allegato = request.FILES.get('file_massivo')
                cf_lista = cf_da_file(allegato) if allegato else None
            if cf_lista is None:
                error = 'Il file non è un CSV o XLSX'
            elif not cf_lista or cf_lista == ['']:
                error = 'Inserisci almeno un codice fiscale.'
            else:
                nome_log = conf['log'][modalita_scelta]
                multi_data = []
                for cf in cf_lista:
                    valido = verifica_cf(cf) in (1, 2)
                    payload = {}
                    if valido:
                        payload, status, purp_id, tok_id = anis_verifica_utente(
                            request.user.username, cf, conf['id_caso'])
                    # invarianti 2 e 8: una riga per interrogazione, solo metadati
                    salva_log(request.user, nome_log, "Verificato utente " + cf,
                              purposeid=purp_id if valido else None,
                              resp_status=status if valido else None,
                              token_id=tok_id if valido else None)
                    messaggio = 'Codice fiscale non corretto' if not valido else _esito_verifica(payload)
                    positivo = messaggio is None
                    risultati.append({
                        'cf': cf, 'titolo': conf['titolo'], 'positivo': positivo,
                        'messaggio': messaggio if not positivo else '',
                        'righe': _righe_servizio(conf['id_caso'], payload) if positivo else [],
                    })
                    # multi_data tiene la forma che le view di export si aspettano: payload nudo
                    # per IFS02/IFS03, (cf, payload) per ANIST, stringa per il CF errato
                    riga = payload if valido else cf + " Codice fiscale non corretto"
                    multi_data.append((cf, riga) if conf['con_cf'] else riga)
                interrogazione = {'titolo': conf['titolo'], 'modalita': modalita_scelta,
                                  'export': conf['export'] if modalita_scelta == 'massiva' else {}}
                if modalita_scelta == 'massiva':
                    request.session['multi_data'] = converti_data(multi_data)

    return render(request, 'istruzione.html', {
        'utente_abilitato': utente_abilitato,
        'servizi': servizi,
        'permessi_json': json.dumps(permessi),
        'risultati': risultati,
        'interrogazione': interrogazione,
        'servizio_scelto': servizio_scelto,
        'modalita_scelta': modalita_scelta,
        'error': error,
    })


def anis_get_voucher(clientid, baseurlauth, client_assertion):
    params = urllib.parse.urlencode({
        'client_id': clientid,
        'client_assertion': client_assertion,
        'client_assertion_type': 'urn:ietf:params:oauth:client-assertion-type:jwt-bearer',
        'grant_type': 'client_credentials'
    })

    headers = {"Content-type": "application/x-www-form-urlencoded"}
    conn = http.client.HTTPSConnection(re.sub(r'^https?://', '', baseurlauth))
    conn.request("POST", "/token.oauth2", params, headers)
    response = conn.getresponse()
    resp_data = response.read()
    voucher_json = json.loads(resp_data)
    voucher = voucher_json.get("access_token")
    if not voucher:
        raise RuntimeError("Voucher non emesso da " + baseurlauth + ": "
                           + str(voucher_json.get("error", resp_data))[:180])
    token_id = None
    try:
        decoded_token = jwt.decode(voucher, options={"verify_signature": False})
        token_id = decoded_token.get('jti')
    except Exception:
        token_id = None
    return voucher, token_id


def anis_verifica_utente(user_ID, cf, id_caso):
    parametri_anis = AnisParametri.objects.get(id=id_caso)
    caso = ''
    if id_caso == 1:
        caso = "IFS02"
    if id_caso == 2:
        caso = "IFS03"
    if id_caso == 3:
        caso = "Anist_Frequenze"
    else:
        caso = "Anist_Titoli"

    kid = parametri_anis.kid
    alg = parametri_anis.alg
    typ = parametri_anis.typ
    issuer = parametri_anis.iss
    subject = parametri_anis.sub
    aud = parametri_anis.aud
    purposeid = parametri_anis.purposeid
    audience = parametri_anis.audience
    baseurlauth = parametri_anis.baseurlauth
    target = parametri_anis.target
    clientid = parametri_anis.clientid
    private_key = parametri_anis.private_key
    userid = user_ID
    location = 'PortaleOpenMSP'
    loa = 'LoA2'
    if id_caso == 1 or id_caso == 2:
        richiesta = f'{{"tax_code":"{cf}"}}'
    else:
        richiesta = f'{{"codiceFiscale":"{cf}"}}'


    issued = datetime.datetime.utcnow()
    delta = datetime.timedelta(minutes=43200)
    expire_in = issued + delta
    dnonce = random.randint(1000000000000, 9999999999999)

    headers_rsa = {
        "kid": kid,
        "alg": alg,
        "typ": typ
    }

    jti = uuid.uuid4()
    audit_payload = {
        "userID": userid,
        "userLocation": location,
        "LoA": loa,
        "iss" : clientid,
        "aud" : audience,
        "purposeId": purposeid,
        "dnonce" : dnonce,
        "jti":str(jti),
        "iat": issued,
        "nbf" : issued,
        "exp": expire_in
        }

    audit = jwt.encode(audit_payload, private_key, algorithm=Algorithms.RS256, headers=headers_rsa)
    audit_hash = hashlib.sha256(audit.encode('UTF-8')).hexdigest()

    jti = uuid.uuid4()
    payload = {
        "iss": clientid,
        "sub": clientid,
        "aud": aud,
        "purposeId": purposeid,
        "jti": str(jti),
        "iat": issued,
        "exp": expire_in,
        "digest": {
            "alg": "SHA256",
            "value": audit_hash
            }
        }

    client_assertion = jwt.encode(payload, private_key, algorithm=Algorithms.RS256, headers=headers_rsa)

    try:
        voucher, token_id = anis_get_voucher(clientid, baseurlauth, client_assertion)
    except Exception as e:
        return {"esito": {"codice": "voucher_error", "descrizione": str(e)[:200]}}, 500, purposeid, None

    # prepara il body per la richiesta e relativo digest
    body = richiesta
    type = 'application/json'
    encoding = 'UTF-8'
    body_digest = hashlib.sha256(body.encode('UTF-8'))
    digest = 'SHA-256=' + base64.b64encode(body_digest.digest()).decode('UTF-8')

    # crea signature
    jti = uuid.uuid4()
    payload = {
        "iss" : clientid,
        "aud" : audience,
        "purposeId": purposeid,
        "sub": clientid,
        "jti": str(jti),
        "iat": issued,
        "nbf" : issued,
        "exp": expire_in,
        "signed_headers": [
            {"digest": digest},
            {"content-type": type},
            {"content-encoding": encoding}
            ]
        }
    signature = jwt.encode(payload, private_key, algorithm=Algorithms.RS256, headers=headers_rsa)
    # effettua chiamata
    api_url = target
    headers =  {"Accept":"application/json",
                "Content-Type":type,
                "Content-Encoding":encoding,
                "Digest":digest,
                "Authorization":"Bearer " + voucher,
                "Agid-JWT-TrackingEvidence":audit,
                "Agid-JWT-Signature":signature
                }

    try:
        response = requests.post(api_url, data=body.encode('UTF-8'), headers=headers, verify=False)
        status_code = response.status_code
        if response.status_code != 200:
            try:
                payload = response.json()
                if isinstance(payload, dict) and (
                    "personal_data" in payload
                    or "enrollments" in payload
                    or "qualifications" in payload
                ):
                    return payload, status_code, purposeid, token_id
            except ValueError:
                pass
            return {
                "esito": {
                    "codice": str(response.status_code),
                    "descrizione": f"Errore API ANIS: {response.text[:200]}"
                }
            }, status_code, purposeid, token_id
        return response.json(), status_code, purposeid, token_id
    except json.JSONDecodeError:
        return {
            "esito": {
                "codice": "JSON_ERR",
                "descrizione": f"Risposta non valida dal server (non JSON): {response.text[:200]}"
            }
        }, 500, purposeid, token_id
    except Exception as e:
        return {
            "esito": {
                "codice": "network_error",
                "descrizione": str(e)
            }
        }, 500, purposeid, token_id

def impostazioni_anis(request):
    servizi_anis = AnisServizi.objects.all()
    parametri_anis = AnisParametri.objects.all()
    svuota_none(parametri_anis)

    service_active = ServiziParametri.objects.all()
    i_serv=0
    service_desc = ["" for _ in range(ServiziParametri.objects.count())]
    for servizio in service_active:
        service_desc[i_serv]= (servizio.attivo)
        i_serv += 1

    if request.method == 'POST':
        active_tab = request.POST.get("active_tab", "").replace("tab", "")

        try:
            active_id = int(active_tab)
        except ValueError:
            active_id = None

        posizione_servizio = list(
            ServiziParametri.objects.filter(gruppo_id=6)
            .order_by('id')
            .values_list('attivo', flat=True)
        )

        ids_to_save = [active_id] if active_id else range(1, AnisParametri.objects.count() + 1)

        for i in ids_to_save:
            if i and i <= len(posizione_servizio) and posizione_servizio[i - 1]:
                kid = request.POST.get('kid' + str(i))
                alg = request.POST.get('alg' + str(i))
                typ = request.POST.get('typ' + str(i))
                iss = request.POST.get('iss' + str(i))
                sub = request.POST.get('sub' + str(i))
                aud = request.POST.get('aud' + str(i))
                purposeid = request.POST.get('purposeid' + str(i))
                audience = request.POST.get('audience' + str(i))
                baseurlauth = request.POST.get('baseurlauth' + str(i))
                target = request.POST.get('target' + str(i))
                clientid = request.POST.get('clientid' + str(i))
                private_key = request.POST.get('private_key' + str(i))
                ver_eservice = request.POST.get('ver_eservice' + str(i))
                dati = AnisParametri(i, i, kid, alg, typ, iss, sub, aud, purposeid, audience, baseurlauth, target, clientid, private_key, ver_eservice)
                dati.save()
        salva_log(request.user,"Impostazioni ANIS", "modifica parametri")

        if active_id:
            return redirect(f"{request.path}#tab{active_id}")
        return redirect(request.path)

    return render(request, 'impostazioni_anis.html', { 'servizi_anis': servizi_anis, 'parametri_anis': parametri_anis })
